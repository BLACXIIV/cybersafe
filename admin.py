import csv
import os
import re
import sqlite3
from datetime import datetime, timedelta, timezone
from functools import wraps
from io import BytesIO, StringIO

import openpyxl
from flask import Blueprint, current_app, flash, g, redirect, render_template, request, send_file, url_for
from werkzeug.utils import secure_filename

from werkzeug.security import generate_password_hash

from auth import login_required, _lookup_by_lrn, _normalize_name
from database.db import get_db
from extensions import limiter
from ranks import BADGE_TITLES, BADGE_ORDER, rank_info
from security import validate_password, describe_problems

bp = Blueprint("admin", __name__, url_prefix="/admin")
ALLOWED_LOGO_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "gif"}

REPORT_PERIODS = {
    "today": "Today",
    "week": "This week",
    "month": "This month",
    "year": "This year",
    "all": "All time",
}

SUSCEPTIBLE_BADGES = ("Bronze", "Silver")

# A voucher holder with no DNS lookup in this window is probably idle, not
# "currently on" the last domain — the live list greys them out instead of
# claiming they're still there.
LIVE_STALE_SECONDS = 300

# "Most visited sites" donut: slice colors are all existing stylesheet
# accents; anything past the top N folds into an "Other" slice in muted grey.
TOP_DOMAIN_LIMIT = 8
TOP_DOMAIN_COLORS = (
    "#22d3ee", "#2f6fed", "#4d8bff", "#22c55e",
    "#f59e0b", "#ef4444", "#facc15", "#a1a1aa",
)
OTHER_DOMAIN_COLOR = "#4b5a78"  # --text-tertiary


def _normalize_grade_name(name):
    """Convert 'Grade 10', '10', 'G10' into 'Grade 10'. Otherwise strip and return as-is."""
    if not name:
        return ""
    name = name.strip()
    digits = ""
    for i, ch in enumerate(name):
        if ch.isdigit():
            digits += ch
        elif digits and not ch.isspace():
            break
    if digits:
        return f"Grade {digits}"
    return name


def _get_or_create_grade(db, grade_name):
    """Return the id of a grade level, creating it if it does not exist."""
    normalized = _normalize_grade_name(grade_name)
    row = db.execute(
        "SELECT id FROM grades WHERE name = ? OR name = ?", (grade_name, normalized)
    ).fetchone()
    if row:
        return row["id"]
    return db.execute("INSERT INTO grades (name) VALUES (?)", (normalized,)).lastrowid


def _get_or_create_level(db, level_number):
    """Return (id, created) for a mission level, inserting a placeholder row
    named 'Mission {n}' when the level_number does not exist yet."""
    row = db.execute(
        "SELECT id FROM levels WHERE level_number = ?", (level_number,)
    ).fetchone()
    if row:
        return row["id"], False
    level_id = db.execute(
        "INSERT INTO levels (level_number, name, focus) VALUES (?, ?, NULL)",
        (level_number, f"Mission {level_number}"),
    ).lastrowid
    return level_id, True


def _insert_student(db, lrn, first_name, last_name, grade_name):
    """Insert a pre-registered student (no password yet). Returns the new id."""
    db.execute(
        """INSERT INTO users
           (full_name, username, email, password_hash, grade_section, role, is_password_set, is_active)
           VALUES (?, ?, ?, '', ?, 'student', 0, 1)""",
        (_normalize_name(first_name, last_name), lrn, f"{lrn}@cybersafe.local", grade_name),
    )
    return db.execute("SELECT id FROM users WHERE username = ?", (lrn,)).fetchone()["id"]


def _import_students_from_excel(db, file_storage):
    """Parse an uploaded .xlsx with columns: LRN, First Name, Last Name, Grade Level.
    Returns (added, skipped, errors)."""
    added, skipped, errors = 0, 0, []

    try:
        wb = openpyxl.load_workbook(BytesIO(file_storage.read()))
        ws = wb.active

        # The provided file has a multi-line title in row 1 and headers in row 2.
        # Try to find the header row.
        header_row = None
        headers = []
        for r in range(1, min(ws.max_row, 10) + 1):
            values = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            lowered = [str(v).lower().strip() if v is not None else "" for v in values]
            if "learner reference number" in lowered or ("first name" in lowered and "last name" in lowered):
                header_row = r
                headers = values
                break

        if header_row is None:
            return 0, 0, ["Could not find a header row with 'Learner Reference Number', 'First Name', 'Last Name', 'Grade Level'."]

        # Map header names to column indexes.
        header_map = {str(h).lower().strip(): i for i, h in enumerate(headers)}

        def col(name):
            return header_map.get(name)

        lrn_idx = col("learner reference number")
        first_idx = col("first name")
        last_idx = col("last name")
        grade_idx = col("grade level")

        if lrn_idx is None or first_idx is None or last_idx is None or grade_idx is None:
            return 0, 0, ["Required columns not found: Learner Reference Number, First Name, Last Name, Grade Level."]

        for r in range(header_row + 1, ws.max_row + 1):
            lrn = ws.cell(row=r, column=lrn_idx + 1).value
            if lrn is None:
                continue
            lrn = str(lrn).strip()
            if not lrn:
                continue

            first_name = str(ws.cell(row=r, column=first_idx + 1).value or "").strip()
            last_name = str(ws.cell(row=r, column=last_idx + 1).value or "").strip()
            grade_name = str(ws.cell(row=r, column=grade_idx + 1).value or "").strip()

            if not first_name or not last_name or not grade_name:
                errors.append(f"Row {r}: missing first name, last name, or grade.")
                continue

            existing = _lookup_by_lrn(db, lrn)
            if existing:
                skipped += 1
                continue

            try:
                _get_or_create_grade(db, grade_name)
                _insert_student(db, lrn, first_name, last_name, grade_name)
                added += 1
            except sqlite3.IntegrityError as e:
                errors.append(f"Row {r}: could not add {lrn} ({e}).")

    except Exception as e:
        errors.append(f"Could not read Excel file: {e}")

    return added, skipped, errors


def _whole_number(value):
    """Return int(value) when it is a whole number, else None."""
    try:
        number = float(str(value).strip())
        if number != int(number):
            return None
        return int(number)
    except (TypeError, ValueError, OverflowError):
        return None


def _import_questions_from_excel(db, file_storage):
    """Parse an uploaded .xlsx with columns: Mission Number, Question Number,
    Question, Choice A–D with Points, Explanation. Missing missions are
    auto-created with a placeholder name.
    Returns (added, skipped, errors, created_missions)."""
    added, skipped, errors, created_missions = 0, 0, [], set()

    try:
        wb = openpyxl.load_workbook(BytesIO(file_storage.read()))
        ws = wb.active

        # Try to find the header row within the first 10 rows.
        header_row = None
        headers = []
        for r in range(1, min(ws.max_row, 10) + 1):
            values = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            lowered = [str(v).lower().strip() if v is not None else "" for v in values]
            if "mission number" in lowered and "question number" in lowered:
                header_row = r
                headers = values
                break

        if header_row is None:
            return 0, 0, ["Could not find a header row with 'Mission Number', 'Question Number', 'Question', 'Choice A'–'Choice D' + Points, 'Explanation'."], []

        # Map header names to column indexes.
        header_map = {str(h).lower().strip(): i for i, h in enumerate(headers)}

        def col(name):
            return header_map.get(name)

        mission_idx = col("mission number")
        qnum_idx = col("question number")
        prompt_idx = col("question")
        explanation_idx = col("explanation")
        choice_idx = {letter: col(f"choice {letter.lower()}") for letter in "ABCD"}
        points_idx = {letter: col(f"choice {letter.lower()} points") for letter in "ABCD"}

        if (mission_idx is None or qnum_idx is None or prompt_idx is None
                or explanation_idx is None
                or any(i is None for i in choice_idx.values())
                or any(i is None for i in points_idx.values())):
            return 0, 0, ["Required columns not found: Mission Number, Question Number, Question, Choice A–D + Points, Explanation."], []

        for r in range(header_row + 1, ws.max_row + 1):
            mission_raw = ws.cell(row=r, column=mission_idx + 1).value
            if mission_raw is None:
                continue
            mission_str = str(mission_raw).strip()
            if not mission_str:
                continue

            qnum_str = str(ws.cell(row=r, column=qnum_idx + 1).value or "").strip()
            prompt = str(ws.cell(row=r, column=prompt_idx + 1).value or "").strip()
            explanation = str(ws.cell(row=r, column=explanation_idx + 1).value or "").strip()
            choice_texts = {
                letter: str(ws.cell(row=r, column=choice_idx[letter] + 1).value or "").strip()
                for letter in "ABCD"
            }

            if not qnum_str or not prompt or any(not text for text in choice_texts.values()):
                errors.append(f"Row {r}: missing question number, question text, or a choice.")
                continue

            mission_number = _whole_number(mission_str)
            question_number = _whole_number(qnum_str)
            if mission_number is None or question_number is None:
                errors.append(f"Row {r}: mission number and question number must be whole numbers.")
                continue

            points = {}
            for letter in "ABCD":
                value = _whole_number(ws.cell(row=r, column=points_idx[letter] + 1).value)
                if value is None or not 0 <= value <= 100:
                    errors.append(f"Row {r}: choice points must be whole numbers between 0 and 100.")
                    points = None
                    break
                points[letter] = value
            if points is None:
                continue

            level_id, level_created = _get_or_create_level(db, mission_number)
            if level_created:
                created_missions.add(mission_number)

            existing = db.execute(
                "SELECT id FROM questions WHERE level_id = ? AND question_number = ?",
                (level_id, question_number),
            ).fetchone()
            if existing:
                skipped += 1
                continue

            try:
                db.execute("SAVEPOINT question_row")
                question_id = db.execute(
                    "INSERT INTO questions (level_id, question_number, prompt, explanation) VALUES (?, ?, ?, ?)",
                    (level_id, question_number, prompt, explanation or None),
                ).lastrowid
                for letter in "ABCD":
                    db.execute(
                        "INSERT INTO choices (question_id, letter, choice_text, points) VALUES (?, ?, ?, ?)",
                        (question_id, letter, choice_texts[letter], points[letter]),
                    )
                db.execute("RELEASE question_row")
                added += 1
                if 100 not in points.values():
                    errors.append(f"Row {r}: no choice scored 100 points; question has no clear correct answer.")
            except sqlite3.IntegrityError as e:
                db.execute("ROLLBACK TO question_row")
                db.execute("RELEASE question_row")
                errors.append(f"Row {r}: could not add question ({e}).")

    except Exception as e:
        errors.append(f"Could not read Excel file: {e}")

    return added, skipped, errors, sorted(created_missions)


def admin_required(view):
    @wraps(view)
    @login_required
    def wrapped_view(*args, **kwargs):
        if g.user["role"] != "admin":
            flash("Administrator access is required.", "error")
            return redirect(url_for("main.dashboard"))
        return view(*args, **kwargs)
    return wrapped_view


def _grade_sort_key(g):
    m = re.search(r"(\d+)", g["name"])
    if m:
        return (int(m.group(1)), g["name"].lower())
    return (0, g["name"].lower())


def _sorted_grades(db):
    return sorted(
        db.execute("SELECT id, name FROM grades").fetchall(),
        key=_grade_sort_key,
    )


def _student_rows(db, max_points):
    """All students decorated with badge/title, sorted highest rank first."""
    students = db.execute(
        """SELECT id, full_name, username, grade_section, points, level, is_active,
                  is_password_set, created_at
           FROM users WHERE role = 'student'
           ORDER BY points DESC, full_name COLLATE NOCASE ASC"""
    ).fetchall()
    rows = []
    for s in students:
        current_badge, current_title, _, _, _, _ = rank_info(s["points"], max_points)
        row = dict(s)
        row["current_badge"] = current_badge
        row["current_title"] = current_title
        rows.append(row)
    rows.sort(key=lambda r: (-BADGE_ORDER[r["current_badge"]], -r["points"]))
    return rows


def _paginate(items, page, per_page, url_builder):
    """Slice `items` for `page` and return (page_items, pagination_dict)."""
    total = len(items)
    total_pages = (total + per_page - 1) // per_page if total else 1
    page = max(1, page)
    if page > total_pages and total:
        page = total_pages
    start = (page - 1) * per_page
    return items[start : start + per_page], {
        "page": page,
        "per_page": per_page,
        "total": total,
        "total_pages": total_pages,
        "has_prev": page > 1,
        "has_next": page < total_pages,
        "prev_url": url_builder(page - 1) if page > 1 else None,
        "next_url": url_builder(page + 1) if page < total_pages else None,
        "start_rank": start,
    }


def _period_start(period):
    """UTC start datetime for a report period, or None for 'all'."""
    now = datetime.utcnow()
    if period == "today":
        return now.replace(hour=0, minute=0, second=0, microsecond=0)
    if period == "week":
        monday = now - timedelta(days=now.weekday())
        return monday.replace(hour=0, minute=0, second=0, microsecond=0)
    if period == "month":
        return now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    if period == "year":
        return now.replace(month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
    return None


def _utc_to_local(iso):
    """'YYYY-MM-DD HH:MM:SS' UTC -> same format in the server's local timezone."""
    return (
        datetime.strptime(iso, "%Y-%m-%d %H:%M:%S")
        .replace(tzinfo=timezone.utc)
        .astimezone()
        .strftime("%Y-%m-%d %H:%M:%S")
    )


def _live_duration(delta):
    """'<1m' / '12m' / '2h 5m' — observed length of a site session."""
    secs = max(0, int(delta.total_seconds()))
    if secs < 60:
        return "<1m"
    if secs < 3600:
        return f"{secs // 60}m"
    hours, mins = divmod(secs // 60, 60)
    return f"{hours}h {mins}m" if mins else f"{hours}h"


def _local_short(iso):
    """'YYYY-MM-DD HH:MM:SS' UTC -> 'Sep 24, 14:32' in the server's timezone."""
    return (
        datetime.strptime(iso, "%Y-%m-%d %H:%M:%S")
        .replace(tzinfo=timezone.utc)
        .astimezone()
        .strftime("%b %d, %H:%M")
    )


def _local_tz_label():
    """Display label for the server's local offset, e.g. 'UTC+8'."""
    minutes = int(datetime.now(timezone.utc).astimezone().utcoffset().total_seconds() // 60)
    sign = "+" if minutes >= 0 else "-"
    hours, rem = divmod(abs(minutes), 60)
    return f"UTC{sign}{hours}" + (f":{rem:02d}" if rem else "")


def _report_data(db, period):
    """Assemble report metrics. Badge/susceptibility state is cumulative (no
    historical snapshots); the period only scopes activity metrics."""
    start = _period_start(period)
    has_period = start is not None
    ap = (start.strftime("%Y-%m-%d %H:%M:%S"),) if has_period else ()
    cond = "ua.answered_at >= ?" if has_period else "1=1"

    total_questions = db.execute("SELECT COUNT(*) AS c FROM questions").fetchone()["c"]
    max_points = total_questions * 100
    rows = _student_rows(db, max_points)

    total = len(rows)
    sus = [r for r in rows if r["current_badge"] in SUSCEPTIBLE_BADGES]
    sus_count = len(sus)
    non_count = total - sus_count
    sus_pct = round((sus_count / total) * 100, 1) if total else 0
    non_pct = round(100 - sus_pct, 1) if total else 0

    badge_dist = {badge: 0 for badge in BADGE_TITLES}
    for r in rows:
        badge_dist[r["current_badge"]] += 1

    grade_map = {}
    for r in rows:
        name = r["grade_section"] or "No grade"
        entry = grade_map.setdefault(name, {"name": name, "total": 0, "sus": 0, "pts": 0})
        entry["total"] += 1
        entry["pts"] += r["points"]
        if r["current_badge"] in SUSCEPTIBLE_BADGES:
            entry["sus"] += 1
    by_grade = list(grade_map.values())
    for entry in by_grade:
        entry["pct"] = round((entry["sus"] / entry["total"]) * 100, 1)
        entry["avg_pts"] = round(entry["pts"] / entry["total"])
    by_grade.sort(key=lambda entry: (-entry["pct"], entry["name"]))

    never_started = db.execute(
        """SELECT COUNT(*) AS c FROM users u WHERE u.role = 'student'
           AND NOT EXISTS (SELECT 1 FROM user_answers ua WHERE ua.user_id = u.id)"""
    ).fetchone()["c"]
    unactivated = db.execute(
        "SELECT COUNT(*) AS c FROM users WHERE role = 'student' AND is_password_set = 0"
    ).fetchone()["c"]
    suspended = db.execute(
        "SELECT COUNT(*) AS c FROM users WHERE role = 'student' AND is_active = 0"
    ).fetchone()["c"]
    avg_points = round(sum(r["points"] for r in rows) / total) if total else 0

    level_count = db.execute("SELECT COUNT(*) AS c FROM levels").fetchone()["c"]
    mastered_all = 0
    if level_count:
        mastered_all = db.execute(
            """SELECT COUNT(*) AS c FROM users u WHERE u.role = 'student'
               AND (SELECT COUNT(*) FROM user_level_progress p
                    WHERE p.user_id = u.id AND p.status = 'mastered') >= ?""",
            (level_count,),
        ).fetchone()["c"]

    activity = db.execute(
        f"""SELECT COUNT(*) AS answers, COALESCE(SUM(ua.points_earned), 0) AS points,
                   COUNT(DISTINCT ua.user_id) AS students
            FROM user_answers ua WHERE {cond}""",
        ap,
    ).fetchone()

    new_registrations = db.execute(
        "SELECT COUNT(*) AS c FROM users WHERE role = 'student' AND created_at >= ?"
        if has_period else
        "SELECT COUNT(*) AS c FROM users WHERE role = 'student'",
        ap,
    ).fetchone()["c"]

    vouchers_issued = db.execute(
        "SELECT COUNT(*) AS c FROM vouchers WHERE created_at >= ?" if has_period else
        "SELECT COUNT(*) AS c FROM vouchers",
        ap,
    ).fetchone()["c"]
    vouchers_active = db.execute(
        "SELECT COUNT(*) AS c FROM vouchers WHERE expires_at > datetime('now')"
    ).fetchone()["c"]

    period_points = {
        row["user_id"]: row["pts"]
        for row in db.execute(
            f"SELECT ua.user_id, SUM(ua.points_earned) AS pts FROM user_answers ua WHERE {cond} GROUP BY ua.user_id",
            ap,
        ).fetchall()
    }

    by_mission = db.execute(
        f"""SELECT l.level_number, l.name,
                   COUNT(ua.id) AS answers,
                   ROUND(AVG(ua.points_earned), 1) AS avg_pts,
                   SUM(CASE WHEN ua.points_earned = 0 THEN 1 ELSE 0 END) AS misses
            FROM user_answers ua
            JOIN questions q ON ua.question_id = q.id
            JOIN levels l ON q.level_id = l.id
            WHERE {cond}
            GROUP BY l.id
            ORDER BY avg_pts ASC""",
        ap,
    ).fetchall()

    missed_questions = db.execute(
        f"""SELECT q.id, q.prompt, l.level_number,
                   COUNT(ua.id) AS answers,
                   SUM(CASE WHEN ua.points_earned = 0 THEN 1 ELSE 0 END) AS misses,
                   ROUND(AVG(ua.points_earned), 1) AS avg_pts
            FROM user_answers ua
            JOIN questions q ON ua.question_id = q.id
            JOIN levels l ON q.level_id = l.id
            WHERE {cond}
            GROUP BY q.id
            HAVING misses > 0
            ORDER BY misses DESC, avg_pts ASC
            LIMIT 5""",
        ap,
    ).fetchall()

    return {
        "students": rows,
        "susceptible": sus,
        "total": total,
        "sus_count": sus_count,
        "non_count": non_count,
        "sus_pct": sus_pct,
        "non_pct": non_pct,
        "badge_dist": badge_dist,
        "by_grade": by_grade,
        "never_started": never_started,
        "unactivated": unactivated,
        "suspended": suspended,
        "avg_points": avg_points,
        "max_points": max_points,
        "mastered_all": mastered_all,
        "level_count": level_count,
        "activity": activity,
        "new_registrations": new_registrations,
        "vouchers_issued": vouchers_issued,
        "vouchers_active": vouchers_active,
        "period_points": period_points,
        "by_mission": by_mission,
        "missed_questions": missed_questions,
    }


def _top_domains(db, start):
    """Most-visited domains since `start` (None = all time), each annotated
    with its share (pct) and donut color. Domains past TOP_DOMAIN_LIMIT fold
    into an "Other" slice so the chart stays readable."""
    cond = "visited_at >= ?" if start else "1=1"
    params = (start.strftime("%Y-%m-%d %H:%M:%S"),) if start else ()
    rows = db.execute(
        f"""SELECT domain, COUNT(*) AS visits FROM site_visits
            WHERE {cond} GROUP BY domain ORDER BY visits DESC LIMIT ?""",
        (*params, TOP_DOMAIN_LIMIT),
    ).fetchall()
    total = db.execute(
        f"SELECT COUNT(*) AS c FROM site_visits WHERE {cond}", params
    ).fetchone()["c"]

    slices = [dict(r) for r in rows]
    other = total - sum(s["visits"] for s in slices)
    if other:
        slices.append({"domain": "Other", "visits": other})
    for i, s in enumerate(slices):
        s["pct"] = round(s["visits"] / total * 100, 1) if total else 0
        s["color"] = TOP_DOMAIN_COLORS[i] if i < len(TOP_DOMAIN_COLORS) else OTHER_DOMAIN_COLOR
    return slices


def _apply_roster_filters(rows, badge_filter, grade_filter, q):
    if badge_filter in BADGE_ORDER:
        rows = [r for r in rows if r["current_badge"] == badge_filter]
    if grade_filter:
        rows = [r for r in rows if (r["grade_section"] or "") == grade_filter]
    if q:
        rows = [
            r for r in rows
            if q in r["full_name"].lower()
            or q in r["username"].lower()
            or (r["grade_section"] and q in r["grade_section"].lower())
            or q in r["current_title"].lower()
        ]
    return rows


@bp.route("/", methods=("GET", "POST"))
@admin_required
def dashboard():
    """The admin home is the Students screen."""
    return redirect(url_for("admin.students"))


@bp.route("/students", methods=("GET", "POST"))
@admin_required
@limiter.limit("60 per minute")
def students():
    db = get_db()

    if request.method == "POST":
        action = request.form.get("action")
        if action == "add_student":
            lrn = request.form.get("lrn", "").strip()
            first_name = request.form.get("first_name", "").strip()
            last_name = request.form.get("last_name", "").strip()
            grade_name = request.form.get("grade_name", "").strip()

            error = None
            if not lrn:
                error = "LRN is required."
            elif not first_name or not last_name:
                error = "First and last name are required."
            elif not grade_name:
                error = "Grade level is required."
            elif _lookup_by_lrn(db, lrn):
                error = "That LRN is already registered."

            if error is None:
                try:
                    _get_or_create_grade(db, grade_name)
                    _insert_student(db, lrn, first_name, last_name, grade_name)
                    db.commit()
                    flash(f"Registered {first_name} {last_name}.", "success")
                except sqlite3.IntegrityError:
                    flash("That LRN is already registered.", "error")
            else:
                flash(error, "error")
        elif action == "import_students":
            file = request.files.get("student_file")
            if not file or not file.filename:
                flash("Choose an Excel file to import.", "error")
            else:
                added, skipped, errors = _import_students_from_excel(db, file)
                db.commit()
                if added:
                    flash(f"Imported {added} student(s).", "success")
                if skipped:
                    flash(f"Skipped {skipped} duplicate(s).", "info")
                for error in errors[:5]:
                    flash(error, "error")
                if len(errors) > 5:
                    flash(f"...and {len(errors) - 5} more errors.", "error")
        elif action == "toggle_user":
            user_id = request.form.get("user_id", type=int)
            if user_id:
                current = db.execute("SELECT is_active, role FROM users WHERE id = ?", (user_id,)).fetchone()
                if current and current["role"] != "admin":
                    new_state = 0 if current["is_active"] else 1
                    db.execute("UPDATE users SET is_active = ? WHERE id = ?", (new_state, user_id))
                    db.commit()
                else:
                    flash("Cannot suspend an admin account.", "error")
        elif action == "reset_password":
            # Lets an admin hand a forgotten account a new password. The same
            # strength rules as signup apply.
            user_id = request.form.get("user_id", type=int)
            password = request.form.get("new_password", "")
            confirm = request.form.get("confirm_password", "")
            student = db.execute(
                "SELECT id, full_name, username, email, role FROM users WHERE id = ?", (user_id,)
            ).fetchone() if user_id else None

            if student is None or student["role"] == "admin":
                flash("Cannot change the password for that account.", "error")
            elif password != confirm:
                flash("Passwords do not match.", "error")
            else:
                problem = describe_problems(
                    validate_password(
                        password,
                        personal_values=(student["full_name"], student["username"], student["email"]),
                    )
                )
                if problem:
                    flash(problem, "error")
                else:
                    db.execute(
                        "UPDATE users SET password_hash = ?, is_password_set = 1 WHERE id = ?",
                        (generate_password_hash(password), student["id"]),
                    )
                    db.commit()
                    flash(f"Password updated for {student['full_name']}.", "success")
        return redirect(url_for("admin.students"))

    total_questions = db.execute("SELECT COUNT(*) AS c FROM questions").fetchone()["c"]
    max_points = total_questions * 100

    badge_filter = request.args.get("badge", "")
    q = request.args.get("q", "").strip().lower()
    page = request.args.get("page", 1, type=int)

    rankings = _apply_roster_filters(_student_rows(db, max_points), badge_filter, "", q)

    def page_url(page_num):
        args = {}
        if badge_filter:
            args["badge"] = badge_filter
        if q:
            args["q"] = q
        if page_num != 1:
            args["page"] = page_num
        return url_for("admin.students", **args)

    paginated_rankings, pagination = _paginate(rankings, page, 10, page_url)

    return render_template(
        "admin_students.html",
        grades=_sorted_grades(db),
        rankings=paginated_rankings,
        pagination=pagination,
        badges=BADGE_TITLES,
        badge_filter=badge_filter,
        q=q,
    )


@bp.route("/missions", methods=("GET", "POST"))
@admin_required
@limiter.limit("60 per minute")
def missions():
    db = get_db()

    if request.method == "POST":
        if request.form.get("action") == "import_questions":
            file = request.files.get("question_file")
            if not file or not file.filename:
                flash("Choose an Excel file to import.", "error")
            else:
                added, skipped, errors, created = _import_questions_from_excel(db, file)
                db.commit()
                if added:
                    flash(f"Imported {added} question(s).", "success")
                if skipped:
                    flash(f"Skipped {skipped} duplicate(s).", "info")
                for mission_number in created:
                    flash(f"Created new Mission {mission_number} (edit its name/focus in the missions admin page).", "info")
                for error in errors[:5]:
                    flash(error, "error")
                if len(errors) > 5:
                    flash(f"...and {len(errors) - 5} more errors.", "error")
        return redirect(url_for("admin.missions"))

    mission_levels = []
    for lvl in db.execute("SELECT * FROM levels ORDER BY level_number").fetchall():
        items = []
        for question in db.execute(
            "SELECT * FROM questions WHERE level_id = ? ORDER BY question_number", (lvl["id"],)
        ).fetchall():
            choices = db.execute(
                "SELECT * FROM choices WHERE question_id = ? ORDER BY letter", (question["id"],)
            ).fetchall()
            items.append({
                "question": question,
                "choices": choices,
                "correct": next((c for c in choices if c["points"] == 100), None),
            })
        mission_levels.append({"level": lvl, "items": items})

    total_questions = db.execute("SELECT COUNT(*) AS c FROM questions").fetchone()["c"]

    return render_template(
        "admin_missions.html",
        mission_levels=mission_levels,
        total_questions=total_questions,
    )


@bp.route("/settings", methods=("GET", "POST"))
@admin_required
@limiter.limit("60 per minute")
def settings():
    db = get_db()

    if request.method == "POST":
        action = request.form.get("action")
        if action == "branding":
            school_name = request.form.get("school_name", "").strip()
            if not school_name:
                flash("School name is required.", "error")
            else:
                logo = request.files.get("school_logo")
                logo_path = db.execute("SELECT logo_path FROM school_settings WHERE id = 1").fetchone()["logo_path"]
                if logo and logo.filename:
                    extension = logo.filename.rsplit(".", 1)[-1].lower() if "." in logo.filename else ""
                    if extension not in ALLOWED_LOGO_EXTENSIONS:
                        flash("Logo must be PNG, JPG, WEBP, or GIF.", "error")
                        return redirect(url_for("admin.settings"))
                    filename = f"school-logo.{extension}"
                    upload_dir = os.path.join(current_app.static_folder, "uploads")
                    os.makedirs(upload_dir, exist_ok=True)
                    logo.save(os.path.join(upload_dir, secure_filename(filename)))
                    logo_path = f"uploads/{filename}"
                db.execute(
                    "UPDATE school_settings SET school_name = ?, logo_path = ? WHERE id = 1",
                    (school_name, logo_path),
                )
                db.commit()
        elif action == "add_grade":
            name = _normalize_grade_name(request.form.get("grade_name", ""))
            if not name:
                flash("Enter a grade name.", "error")
            else:
                try:
                    db.execute("INSERT INTO grades (name) VALUES (?)", (name,))
                    db.commit()
                except sqlite3.IntegrityError:
                    flash("That grade already exists.", "error")
        elif action == "delete_grade":
            db.execute("DELETE FROM grades WHERE id = ?", (request.form.get("grade_id", type=int),))
            db.commit()
        return redirect(url_for("admin.settings"))

    return render_template(
        "admin_settings.html",
        settings=db.execute("SELECT * FROM school_settings WHERE id = 1").fetchone(),
        grades=_sorted_grades(db),
    )


@bp.route("/download-student-template")
@admin_required
@limiter.limit("10 per minute")
def download_student_template():
    """Return an empty .xlsx template with the columns required for bulk import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LearnerDataEntry"
    headers = ["Learner Reference Number", "First Name", "Last Name", "Grade Level"]
    ws.append(headers)
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col).font = openpyxl.styles.Font(bold=True)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="cybersafe-student-import-template.xlsx",
    )


@bp.route("/download-question-template")
@admin_required
@limiter.limit("10 per minute")
def download_question_template():
    """Return an empty .xlsx template with the columns required for bulk import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QuestionDataEntry"
    headers = [
        "Mission Number", "Question Number", "Question",
        "Choice A", "Choice A Points", "Choice B", "Choice B Points",
        "Choice C", "Choice C Points", "Choice D", "Choice D Points",
        "Explanation",
    ]
    ws.append(headers)
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col).font = openpyxl.styles.Font(bold=True)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="cybersafe-question-import-template.xlsx",
    )


def _report_filters():
    return (
        request.args.get("badge", ""),
        request.args.get("grade", ""),
        request.args.get("q", "").strip().lower(),
    )


def _report_args(current_period, badge_filter, grade_filter, q, **extra):
    args = dict(extra)
    args.setdefault("period", current_period)
    if args["period"] == "all":
        del args["period"]
    if badge_filter:
        args["badge"] = badge_filter
    if grade_filter:
        args["grade"] = grade_filter
    if q:
        args["q"] = q
    return args


@bp.route("/report")
@admin_required
def report_redirect():
    """Legacy URL — the report screen now lives at /admin/analytics."""
    return redirect(url_for("admin.analytics", **request.args))


@bp.route("/analytics")
@admin_required
@limiter.limit("60 per minute")
def analytics():
    """Susceptibility report: roster state plus period-scoped activity metrics."""
    db = get_db()
    period = request.args.get("period", "all")
    if period not in REPORT_PERIODS:
        period = "all"
    badge_filter, grade_filter, q = _report_filters()
    page = request.args.get("page", 1, type=int)

    data = _report_data(db, period)
    for row in data["susceptible"]:
        row["period_pts"] = data["period_points"].get(row["id"], 0)

    filtered = _apply_roster_filters(
        sorted(data["susceptible"], key=lambda r: (r["points"], r["full_name"].lower())),
        badge_filter, grade_filter, q,
    )

    def page_url(page_num):
        args = _report_args(period, badge_filter, grade_filter, q)
        if page_num != 1:
            args["page"] = page_num
        return url_for("admin.analytics", **args)

    paginated, pagination = _paginate(filtered, page, 10, page_url)

    return render_template(
        "admin_analytics.html",
        **{k: v for k, v in data.items() if k != "susceptible"},
        top_domains=_top_domains(db, _period_start(period)),
        period=period,
        periods=REPORT_PERIODS,
        period_label=REPORT_PERIODS[period],
        susceptible=paginated,
        pagination=pagination,
        q=q,
        badge_filter=badge_filter,
        grade_filter=grade_filter,
        grade_options=[g["name"] for g in _sorted_grades(db)],
        badges=BADGE_TITLES,
        report_args=lambda **kw: _report_args(period, badge_filter, grade_filter, q, **kw),
    )


@bp.route("/activity")
@admin_required
@limiter.limit("60 per minute")
def activity():
    """Site visits: domains resolved by devices holding an active voucher,
    joined back to the student who redeemed it. Domain-level only."""
    db = get_db()
    period = request.args.get("period", "all")
    if period not in REPORT_PERIODS:
        period = "all"
    q = request.args.get("q", "").strip().lower()
    page = request.args.get("page", 1, type=int)

    start = _period_start(period)
    cond = "sv.visited_at >= ?" if start else "1=1"
    params = (start.strftime("%Y-%m-%d %H:%M:%S"),) if start else ()
    visits = db.execute(
        f"""SELECT sv.visited_at, sv.domain,
                   u.full_name, u.username, u.grade_section
            FROM site_visits sv
            JOIN users u ON u.id = sv.user_id
            WHERE {cond} AND u.role = 'student'
            ORDER BY sv.visited_at DESC, sv.id DESC""",
        params,
    ).fetchall()

    if q:
        visits = [
            v for v in visits
            if q in v["full_name"].lower()
            or q in v["username"].lower()
            or q in v["domain"].lower()
        ]

    def page_url(page_num):
        args = _report_args(period, "", "", q)
        if page_num != 1:
            args["page"] = page_num
        return url_for("admin.activity", **args)

    paginated, pagination = _paginate(visits, page, 10, page_url)
    paginated = [
        {key: v[key] for key in v.keys()} | {"visited_at": _utc_to_local(v["visited_at"])}
        for v in paginated
    ]

    # "Live now": students holding an active voucher are the ones able to
    # surf, so they're the live set — each shown with their latest session:
    # the domain, how long it has run, and when it started.
    live_rows = db.execute(
        """SELECT u.full_name, u.username, u.grade_section,
                  s.domain AS last_domain, s.started_at, s.last_seen_at
           FROM users u
           LEFT JOIN site_sessions s ON s.id = (
               SELECT s2.id FROM site_sessions s2 WHERE s2.user_id = u.id
               ORDER BY s2.last_seen_at DESC, s2.id DESC LIMIT 1)
           WHERE u.role = 'student' AND EXISTS (
               SELECT 1 FROM vouchers v
               WHERE v.user_id = u.id AND v.used_at IS NOT NULL
                 AND v.expires_at > datetime('now'))
           ORDER BY s.last_seen_at DESC"""
    ).fetchall()

    now = datetime.utcnow()
    live = []
    for row in live_rows:
        if row["last_seen_at"]:
            last_seen = datetime.strptime(row["last_seen_at"], "%Y-%m-%d %H:%M:%S")
            started = datetime.strptime(row["started_at"], "%Y-%m-%d %H:%M:%S")
            age = (now - last_seen).total_seconds()
            ongoing = age < LIVE_STALE_SECONDS
            # Ongoing counts up to now; ended sessions show the observed span.
            duration = _live_duration((now if ongoing else last_seen) - started)
            detail = (
                f"{duration} · ongoing since {_local_short(row['started_at'])}"
                if ongoing else
                f"{duration} · ended {_local_short(row['last_seen_at'])}"
            )
        else:
            ongoing = False
            detail = "no lookups yet"
        live.append({
            "full_name": row["full_name"],
            "username": row["username"],
            "grade_section": row["grade_section"],
            "last_domain": row["last_domain"],
            "detail": detail,
            "stale": not ongoing,
        })

    return render_template(
        "admin_activity.html",
        visits=paginated,
        live=live,
        tz_label=_local_tz_label(),
        pagination=pagination,
        period=period,
        periods=REPORT_PERIODS,
        period_label=REPORT_PERIODS[period],
        q=q,
        report_args=lambda **kw: _report_args(period, "", "", q, **kw),
    )


_STUDENT_EXPORT_HEADERS = [
    "Name", "LRN", "Grade Level", "Points", "Badge", "Rank Title",
    "Points In Period", "Susceptible", "Account Status", "Password Set", "Registered",
]


def _student_export_row(r, period_points):
    return [
        r["full_name"],
        r["username"],
        r["grade_section"] or "",
        r["points"],
        r["current_badge"],
        r["current_title"],
        period_points.get(r["id"], 0),
        "Yes" if r["current_badge"] in SUSCEPTIBLE_BADGES else "No",
        "Suspended" if not r["is_active"] else "Active",
        "Yes" if r["is_password_set"] else "No",
        r["created_at"],
    ]


@bp.route("/analytics/export")
@admin_required
@limiter.limit("10 per minute")
def analytics_export():
    """Download the report as .xlsx (multi-sheet) or .csv (student list)."""
    fmt = request.args.get("format", "xlsx")
    period = request.args.get("period", "all")
    if period not in REPORT_PERIODS:
        period = "all"
    badge_filter, grade_filter, q = _report_filters()

    db = get_db()
    data = _report_data(db, period)
    rows = _apply_roster_filters(data["students"], badge_filter, grade_filter, q)
    stamp = datetime.utcnow().strftime("%Y%m%d")

    if fmt == "csv":
        text = StringIO()
        writer = csv.writer(text)
        writer.writerow(_STUDENT_EXPORT_HEADERS)
        for r in rows:
            writer.writerow(_student_export_row(r, data["period_points"]))
        buffer = BytesIO(text.getvalue().encode("utf-8-sig"))
        return send_file(
            buffer,
            mimetype="text/csv",
            as_attachment=True,
            download_name=f"cybersafe-report-{period}-{stamp}.csv",
        )

    wb = openpyxl.Workbook()
    bold = openpyxl.styles.Font(bold=True)

    summary = wb.active
    summary.title = "Summary"
    school_name = db.execute("SELECT school_name FROM school_settings WHERE id = 1").fetchone()["school_name"]
    summary_rows = [
        ("School", school_name),
        ("Generated (UTC)", datetime.utcnow().strftime("%Y-%m-%d %H:%M")),
        ("Period", REPORT_PERIODS[period]),
        ("", ""),
        ("Students enrolled", data["total"]),
        ("Susceptible (Bronze/Silver)", f"{data['sus_count']} ({data['sus_pct']}%)"),
        ("At or above Gold", f"{data['non_count']} ({data['non_pct']}%)"),
        ("Average points", data["avg_points"]),
        ("Never started a mission", data["never_started"]),
        ("Password not set", data["unactivated"]),
        ("Suspended", data["suspended"]),
        ("Mastered all missions", data["mastered_all"]),
        ("", ""),
        (f"Answers ({REPORT_PERIODS[period]})", data["activity"]["answers"]),
        (f"Points earned ({REPORT_PERIODS[period]})", data["activity"]["points"]),
        (f"Active students ({REPORT_PERIODS[period]})", data["activity"]["students"]),
        (f"New registrations ({REPORT_PERIODS[period]})", data["new_registrations"]),
        (f"Vouchers issued ({REPORT_PERIODS[period]})", data["vouchers_issued"]),
        ("Vouchers active now", data["vouchers_active"]),
        ("", ""),
        ("Badge distribution", ""),
    ] + [(f"  {badge} · {BADGE_TITLES[badge]}", data["badge_dist"][badge]) for badge in BADGE_TITLES]
    for label, value in summary_rows:
        summary.append([label, value])
    summary["A1"].font = bold
    for cell in summary["A"]:
        if cell.value in ("Badge distribution",):
            cell.font = bold

    students_ws = wb.create_sheet("Students")
    students_ws.append(_STUDENT_EXPORT_HEADERS)
    for cell in students_ws[1]:
        cell.font = bold
    for r in rows:
        students_ws.append(_student_export_row(r, data["period_points"]))

    grades_ws = wb.create_sheet("By Grade")
    grades_ws.append(["Grade Level", "Students", "Susceptible", "Susceptible %", "Avg Points"])
    for cell in grades_ws[1]:
        cell.font = bold
    for g_row in data["by_grade"]:
        grades_ws.append([g_row["name"], g_row["total"], g_row["sus"], g_row["pct"], g_row["avg_pts"]])

    missions_ws = wb.create_sheet("By Mission")
    missions_ws.append(["Mission", "Name", "Answers", "Avg Points", "Zero-Point Answers"])
    for cell in missions_ws[1]:
        cell.font = bold
    for m in data["by_mission"]:
        missions_ws.append([f"Mission {m['level_number']}", m["name"], m["answers"], m["avg_pts"], m["misses"]])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"cybersafe-report-{period}-{stamp}.xlsx",
    )
